import io
import re
import hashlib
import unicodedata
from typing import List, Dict, Any, Tuple, Optional
import openpyxl
from fastapi import HTTPException, status
from sqlmodel.ext.asyncio.session import AsyncSession
from sqlmodel import select

from backend.app.models.entities import (
    Rider, Team, Category, StageEntry, RaceStatusEnum, PaymentStatusEnum, GenderEnum
)
from backend.app.services.scoring import (
    recalculate_stage_ranks, recalculate_championship_standings, recalculate_team_standings,
    merge_duplicate_entries
)
from backend.app.services.pdf_service import (
    parse_time_to_ms, format_ms_to_time, collapse_team_hyphen,
)


def slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_-]+", "-", text)
    return text.strip("-")


def normalize_name(name: str) -> str:
    """
    Normaliza un nombre de ciclista para identidad determinística:
    quita acentos/diacríticos (PEÑA -> PENA, LÓPEZ -> LOPEZ), pasa a mayúsculas,
    elimina puntuación y colapsa espacios. Evita colisiones por tildes y sufijos.
    """
    text = unicodedata.normalize("NFKD", name or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"[^A-Z0-9 ]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def generate_deterministic_dni(name: str) -> str:
    """Genera un DNI sintético numérico de 8 dígitos basado en el hash del nombre normalizado."""
    clean = normalize_name(name).replace(" ", "")
    h = hashlib.sha256(clean.encode("utf-8")).hexdigest()
    num = int(h[:8], 16) % 90000000 + 10000000
    return str(num)


async def resolve_category(session: AsyncSession, raw_cat_name: str) -> Optional[Category]:
    """Resuelve la categoría canónica comparando nombre, short_code y aliases."""
    clean_cat = raw_cat_name.strip().upper()
    
    # 1. Búsqueda directa
    stmt = select(Category).where(Category.name.ilike(clean_cat))
    res = await session.exec(stmt)
    cat = res.first()
    if cat:
        return cat

    # 2. Búsqueda por short_code
    stmt2 = select(Category).where(Category.short_code.ilike(clean_cat))
    res2 = await session.exec(stmt2)
    cat2 = res2.first()
    if cat2:
        return cat2

    # 3. Búsqueda en aliases
    stmt3 = select(Category)
    res3 = await session.exec(stmt3)
    all_cats = res3.all()
    for c in all_cats:
        for alias in (c.aliases or []):
            if alias.upper() in clean_cat or clean_cat in alias.upper():
                return c

    return None


async def resolve_or_create_team(session: AsyncSession, raw_team_name: Optional[str]) -> Optional[int]:
    """
    Normaliza nombres de equipo y retorna team_id o None si es Libre/Independiente.

    El PDF oficial parte algunos nombres de club entre dos líneas
    ("CLUB ORIGINAL BIKES AREQUIPA-" / "PERÚ"), lo que deja un guion pegado a un
    espacio. Sin normalizar, ese club se crearía duplicado del homónimo limpio y
    sus puntos se partirían en dos equipos distintos.
    """
    if not raw_team_name:
        return None
    cleaned = collapse_team_hyphen(raw_team_name.strip().upper())
    cleaned = re.sub(r"\s+", " ", cleaned)
    if cleaned in {"LIBRE", "INDEPENDIENTE", "TEAM", "-", "NINGUNO", ""}:
        return None

    stmt = select(Team).where(Team.name.ilike(cleaned))
    res = await session.exec(stmt)
    team = res.first()
    if team:
        return team.id

    # Crear equipo si no existe
    slug = slugify(cleaned)
    new_team = Team(name=cleaned, slug=slug, is_active=True)
    session.add(new_team)
    await session.flush()
    return new_team.id


async def import_riders_from_excel(
    session: AsyncSession,
    stage_id: int,
    file_bytes: bytes,
    is_legacy_format: bool = False,
    merge_duplicates: bool = True
) -> Dict[str, Any]:
    """
    Importación transaccional con validación atómica (Dry-Run implícito).
    Si hay al menos 1 error crítico, arroja HTTP 422 con lista de errores.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet_name = "LISTA" if "LISTA" in wb.sheetnames else ("Inscritos" if "Inscritos" in wb.sheetnames else wb.sheetnames[0])
    ws = wb[sheet_name]

    errors = []
    rows_data = []
    used_bibs = set()

    current_category_context = None

    for row_idx in range(1, ws.max_row + 1):
        vals = [ws.cell(row_idx, col_idx).value for col_idx in range(1, ws.max_column + 1)]
        if not any(v is not None for v in vals):
            continue

        str_vals = [str(v).strip() if v is not None else "" for v in vals]
        first_cell = str_vals[0].upper()

        # Detección de fila de cabecera
        if "DORSAL" in first_cell or "DNI" in first_cell or "INSCRITOS" in first_cell:
            # En el formato de LA JOYA EXTREME, la columna 4 contiene el nombre de la categoría del bloque
            if len(vals) >= 4 and vals[3] and str(vals[3]).strip():
                header_candidate = str(vals[3]).strip()
                if "TOTAL" not in header_candidate.upper() and "PUNTOS" not in header_candidate.upper():
                    current_category_context = header_candidate
            continue

        # Lectura de datos
        # Caso A: Legacy LA JOYA (DORSAL, NOMBRE, TEAM, CATEGORIA/SUBCAT, T.TOTAL, PUNTOS, PUESTO)
        # Caso B: Oficial (DNI, NOMBRE_COMPLETO, GENERO, TEAM, CATEGORIA, DORSAL, TIEMPO, ESTADO_PAGO)
        raw_bib = None
        raw_dni = None
        raw_name = None
        raw_team = None
        raw_cat = None
        raw_time = None
        raw_gender = None

        if is_legacy_format or ("DORSAL" in [str(ws.cell(r, 1).value).upper() for r in range(1, min(5, ws.max_row)) if ws.cell(r, 1).value]):
            # Fila tipo LA JOYA: Col 1 = Dorsal, Col 2 = Nombre, Col 3 = Team, Col 4 = Cat
            try:
                raw_bib = int(vals[0]) if vals[0] is not None and str(vals[0]).isdigit() else None
            except Exception:
                raw_bib = None

            raw_name = str_vals[1] if len(str_vals) > 1 else ""
            raw_team = str_vals[2] if len(str_vals) > 2 else ""
            raw_cat = str_vals[3] if len(str_vals) > 3 and str_vals[3] else current_category_context
            raw_time = str_vals[4] if len(str_vals) > 4 and str_vals[4] else None

            if not raw_name:
                continue

            raw_dni = generate_deterministic_dni(raw_name)
            raw_gender = GenderEnum.F if "DAMAS" in (raw_cat or "").upper() else GenderEnum.M
        else:
            # Formato estándar de especificación
            raw_dni = str_vals[0] if len(str_vals) > 0 else ""
            raw_name = str_vals[1] if len(str_vals) > 1 else ""
            raw_gender = GenderEnum.F if (len(str_vals) > 2 and str_vals[2].upper().startswith("F")) else GenderEnum.M
            raw_team = str_vals[3] if len(str_vals) > 3 else ""
            raw_cat = str_vals[4] if len(str_vals) > 4 else ""
            raw_bib_val = str_vals[5] if len(str_vals) > 5 else ""
            raw_bib = int(raw_bib_val) if raw_bib_val.isdigit() else None
            raw_time = str_vals[6] if len(str_vals) > 6 and str_vals[6] else None

        # Validaciones de integridad
        if not raw_name or len(raw_name) < 3:
            errors.append({"row": row_idx, "field": "nombre", "error": "Nombre del ciclista requerido", "raw_value": raw_name})
            continue

        if not raw_dni:
            errors.append({"row": row_idx, "field": "dni", "error": "DNI obligatorio", "raw_value": ""})
            continue

        if raw_bib is not None:
            if raw_bib in used_bibs:
                errors.append({"row": row_idx, "field": "bib_number", "error": f"Dorsal {raw_bib} duplicado en esta fecha", "raw_value": str(raw_bib)})
                continue
            used_bibs.add(raw_bib)

        # Resolver categoría
        cat_obj = await resolve_category(session, raw_cat or "")
        if not cat_obj:
            errors.append({"row": row_idx, "field": "categoria", "error": f"Categoría '{raw_cat}' no reconocida", "raw_value": raw_cat})
            continue

        rows_data.append({
            "dni": raw_dni,
            "full_name": re.sub(r"[\s\-.;]+$", "", raw_name).strip(),
            "gender": raw_gender,
            "team_name": raw_team,
            "category_id": cat_obj.id,
            "bib_number": raw_bib or 0,
            "time_str": raw_time
        })

    # Si hay errores críticos, abortar con HTTP 422
    if errors:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"message": f"Se encontraron {len(errors)} errores en el archivo Excel.", "errors": errors}
        )

    # Si 0 errores, procesar UPSERT
    imported_count = 0
    for r in rows_data:
        team_id = await resolve_or_create_team(session, r["team_name"])

        # Buscar o crear Rider
        stmt_rider = select(Rider).where(Rider.dni == r["dni"])
        res_rider = await session.exec(stmt_rider)
        rider = res_rider.first()

        if not rider:
            rider = Rider(
                dni=r["dni"],
                full_name=r["full_name"],
                slug=f"{slugify(r['full_name'])}-{r['dni'][-4:]}",
                team_id=team_id,
                gender=r["gender"]
            )
            session.add(rider)
            await session.flush()
        else:
            rider.full_name = r["full_name"]
            rider.team_id = team_id
            session.add(rider)

        # Buscar o crear StageEntry.
        # La clave es (stage_id, bib_number): así un mismo ciclista (mismo DNI) con dos
        # dorsales en la misma fecha NO se colapsa durante la importación; la fusión por
        # identidad se resuelve después (conservando el mejor tiempo) en merge_duplicate_entries.
        if r["bib_number"]:
            stmt_entry = select(StageEntry).where(
                StageEntry.stage_id == stage_id,
                StageEntry.bib_number == r["bib_number"]
            )
        else:
            stmt_entry = select(StageEntry).where(
                StageEntry.stage_id == stage_id,
                StageEntry.rider_id == rider.id
            )
        res_entry = await session.exec(stmt_entry)
        entry = res_entry.first()

        elapsed_ms = parse_time_to_ms(r["time_str"]) if r["time_str"] else None
        race_status = RaceStatusEnum.FINISHED if elapsed_ms else RaceStatusEnum.REGISTERED

        if not entry:
            entry = StageEntry(
                stage_id=stage_id,
                rider_id=rider.id,
                category_id=r["category_id"],
                team_id=team_id,
                bib_number=r["bib_number"],
                payment_status=PaymentStatusEnum.PAID,
                race_status=race_status,
                finish_elapsed_ms=elapsed_ms
            )
        else:
            entry.category_id = r["category_id"]
            entry.team_id = team_id
            entry.bib_number = r["bib_number"]
            if elapsed_ms:
                entry.finish_elapsed_ms = elapsed_ms
                entry.race_status = race_status

        session.add(entry)
        imported_count += 1

    await session.commit()

    # Fusión de identidades duplicadas dentro de la misma fecha (mismo rider, varios dorsales)
    if merge_duplicates:
        await merge_duplicate_entries(session, stage_id)

    # Recalcular puestos y rankings
    await recalculate_stage_ranks(session, stage_id)
    await recalculate_championship_standings(session)
    await recalculate_team_standings(session, stage_id)

    return {
        "status": "success",
        "imported_riders": imported_count,
        "stage_id": stage_id
    }


async def export_stage_results_to_excel(session: AsyncSession, stage_id: int) -> bytes:
    """Genera en memoria un libro .xlsx con las 3 hojas reglamentarias.

    Cada hoja se resuelve con UNA consulta: corredor, categoría y equipo
    salen con `JOIN` en la misma sentencia. Antes cada fila recargaba sus
    relaciones con un `refresh` (una consulta por fila y por hoja).
    """
    wb = openpyxl.Workbook()
    # Hoja 1: Resultados_Fecha
    ws1 = wb.active
    ws1.title = "Resultados_Fecha"
    ws1.append(["PUESTO", "DORSAL", "CICLISTA", "DNI", "TEAM", "CATEGORIA", "TIEMPO OFICIAL", "PUNTOS"])

    stmt1 = (
        select(StageEntry, Rider, Category, Team)
        .outerjoin(Rider, StageEntry.rider_id == Rider.id)
        .outerjoin(Category, StageEntry.category_id == Category.id)
        .outerjoin(Team, StageEntry.team_id == Team.id)
        .where(StageEntry.stage_id == stage_id)
        .order_by(StageEntry.category_id, StageEntry.category_rank)
    )
    res1 = await session.exec(stmt1)
    for entry, rider, category, team in res1.all():
        ws1.append([
            entry.category_rank or "-",
            entry.bib_number,
            rider.full_name if rider else "",
            rider.dni if rider else "",
            team.name if team else "LIBRE",
            category.name if category else "",
            format_ms_to_time(entry.finish_elapsed_ms or 0) if entry.finish_elapsed_ms else "DNF/DNS",
            entry.points_earned
        ])

    # Hoja 2: Ranking_Acumulado_Individual
    ws2 = wb.create_sheet("Ranking_Acumulado_Individual")
    ws2.append(["PUESTO", "DELTA", "CICLISTA", "CATEGORIA", "TEAM", "TOTAL PUNTOS", "VICTORIAS", "TIEMPO ACUMULADO"])

    from backend.app.models.entities import ChampionshipStanding
    # El club de esta hoja es el del CORREDOR (no el de la inscripción).
    stmt_stand = (
        select(ChampionshipStanding, Rider, Category, Team)
        .outerjoin(Rider, ChampionshipStanding.rider_id == Rider.id)
        .outerjoin(Category, ChampionshipStanding.category_id == Category.id)
        .outerjoin(Team, Rider.team_id == Team.id)
        .order_by(ChampionshipStanding.category_id, ChampionshipStanding.current_rank)
    )
    res2 = await session.exec(stmt_stand)
    for st, rider, cat, team in res2.all():
        delta_str = f"▲ +{st.rank_delta}" if st.rank_delta > 0 else (f"▼ {st.rank_delta}" if st.rank_delta < 0 else "=")
        ws2.append([
            st.current_rank,
            delta_str,
            rider.full_name if rider else "",
            cat.name if cat else "",
            team.name if team else "LIBRE",
            st.total_points,
            st.first_places_count,
            format_ms_to_time(st.total_elapsed_ms)
        ])

    # Hoja 3: Ranking_Equipos
    ws3 = wb.create_sheet("Ranking_Equipos")
    ws3.append(["PUESTO", "DELTA", "EQUIPO", "TOTAL PUNTOS", "OROS (1°)", "PLATAS (2°)", "BRONCES (3°)", "CORREDORES"])
    from backend.app.models.entities import TeamChampionshipStanding
    stmt_team = (
        select(TeamChampionshipStanding, Team)
        .outerjoin(Team, TeamChampionshipStanding.team_id == Team.id)
        .order_by(TeamChampionshipStanding.current_rank)
    )
    res3 = await session.exec(stmt_team)
    for ts, team in res3.all():
        delta_str = f"▲ +{ts.rank_delta}" if ts.rank_delta > 0 else (f"▼ {ts.rank_delta}" if ts.rank_delta < 0 else "=")
        ws3.append([
            ts.current_rank,
            delta_str,
            team.name if team else "",
            ts.total_points,
            ts.first_places_count,
            ts.second_places_count,
            ts.third_places_count,
            ts.total_riders_participations
        ])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
