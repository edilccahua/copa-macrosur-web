import io
import datetime
import pytest
import openpyxl
from fastapi import HTTPException
from sqlmodel.ext.asyncio.session import AsyncSession
from sqlmodel import select

from backend.app.models.entities import Stage, Category, Rider, StageEntry
from backend.app.services.excel_service import import_riders_from_excel, export_stage_results_to_excel
from backend.app.services.excel_service import generate_deterministic_dni, normalize_name


def test_normalize_name_removes_accents_and_punctuation():
    assert normalize_name("JOSÉ PEÑA.") == "JOSE PENA"
    assert normalize_name("LÓPEZ-QUISPE") == "LOPEZ QUISPE"
    # El DNI sintético es estable ante tildes/puntuación
    assert generate_deterministic_dni("JOSE PEÑA.") == generate_deterministic_dni("JOSE PENA")
    assert generate_deterministic_dni("ROBERTO VALENCIA") != generate_deterministic_dni("ROBERTO VALENCIA GONZALES")


def create_mock_excel(rows: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LISTA"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_excel_dry_run_rollback_on_duplicate_bib(async_session: AsyncSession):
    stage = Stage(stage_number=1, name="Fecha 1", slug="f1", date=datetime.date(2026, 3, 15))
    cat = Category(name="ELITE", short_code="ELI", aliases=["ELITE"])
    async_session.add_all([stage, cat])
    await async_session.commit()

    # Excel con dos filas con el mismo dorsal (15)
    excel_data = [
        ["DORSAL", "NOMBRE", "TEAM", "ELITE", "T. TOTAL", "PUNTOS", "PUESTO"],
        [15, "JUAN PEREZ", "LIBRE", "ELITE", None, None, None],
        [15, "PEDRO GOMEZ", "LIBRE", "ELITE", None, None, None]  # DORSAL DUPLICADO
    ]
    excel_bytes = create_mock_excel(excel_data)

    # Debe lanzar HTTPException 422
    with pytest.raises(HTTPException) as exc_info:
        await import_riders_from_excel(async_session, stage.id, excel_bytes, is_legacy_format=True)

    assert exc_info.value.status_code == 422
    assert "duplicado" in str(exc_info.value.detail).lower()

    # Verificar que NO se escribió ningún ciclista en BD (Rollback total)
    riders = (await async_session.exec(select(Rider))).all()
    assert len(riders) == 0


@pytest.mark.asyncio
async def test_excel_export_three_sheets(async_session: AsyncSession):
    stage = Stage(stage_number=1, name="Fecha 1", slug="f1", date=datetime.date(2026, 3, 15))
    cat = Category(name="ELITE", short_code="ELI", aliases=["ELITE"])
    async_session.add_all([stage, cat])
    await async_session.commit()

    # Importar un corredor válido
    excel_data = [
        ["DORSAL", "NOMBRE", "TEAM", "ELITE", "T. TOTAL", "PUNTOS", "PUESTO"],
        [7, "OSCAR CARDENAS", "LIBRE", "ELITE", "01:15:30", None, None]
    ]
    excel_bytes = create_mock_excel(excel_data)
    await import_riders_from_excel(async_session, stage.id, excel_bytes, is_legacy_format=True)

    # Exportar
    exported_bytes = await export_stage_results_to_excel(async_session, stage.id)
    wb = openpyxl.load_workbook(io.BytesIO(exported_bytes))

    # Verificar que existan las 3 hojas reglamentarias
    assert "Resultados_Fecha" in wb.sheetnames
    assert "Ranking_Acumulado_Individual" in wb.sheetnames
    assert "Ranking_Equipos" in wb.sheetnames

    ws1 = wb["Resultados_Fecha"]
    assert ws1.cell(2, 3).value == "OSCAR CARDENAS"
